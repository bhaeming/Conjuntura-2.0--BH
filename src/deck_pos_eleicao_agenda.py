"""
Seção 06 (Agenda 2027) do deck "Pós-eleição 2027 v4", enxuta e visual:

- "O que acompanhar" e "Balanço de riscos" viram um só slide: um mapa em que cada tema traz,
  no centro, o indicador a acompanhar, com os riscos para pior à esquerda e para melhor à direita;
- "Mensagens finais" vira quatro mensagens numeradas sobre a reta de chegada da estrada do roteiro.

Edita o deck já ajustado à mão (só esses slides, com backup):

    python src/deck_pos_eleicao_agenda.py
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_pos_eleicao_roteiro import BLUE, BRANCO, INK, PISTA, _sem_estilo, _texto  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs_provisorio" / "Pós-eleição 2027 - Temas econômicos (versão final).pptx"
MUTED, TINTA_AZUL = RGBColor(0x6B, 0x72, 0x80), RGBColor(0xE6, 0xEC, 0xFA)

# tema, indicador a acompanhar, risco para pior, risco para melhor
RISCOS = [
    ("Fiscal", "Resultado primário e dívida bruta",
     "Revisão da meta ou adiamento do ajuste elevam o prêmio de risco",
     "Ajuste crível nas despesas obrigatórias acelera a queda da Selic"),
    ("Juros e inflação", "Focus e comunicação do Copom",
     "Novas altas do Fed e juro de 10 anos dos EUA em 5%",
     "Atividade mais fraca alivia os preços de serviços"),
    ("Famílias e crédito", "Comprometimento de renda",
     "Inadimplência recorde restringe a oferta de crédito",
     "Renegociação de dívidas e fim das apostas liberam renda"),
    ("Cenário externo", "Brent e câmbio",
     "Petróleo, geopolítica e El Niño pressionam preços, fretes e safra",
     "Petróleo e agro sustentam câmbio e receitas"),
    ("Investimento", "FBCF e produção industrial",
     "Crédito caro adia projetos e compra de bens de capital",
     "Juros menores destravam FBCF, construção e bens de capital"),
]

MENSAGENS = [
    ("Crescimento menor", "PIB do Brasil entre 1,3% e 1,7% em 2027, com desemprego mínimo. "
                          "SC segue acima da média nacional."),
    ("Fiscal destrava juros", "Com inflação acima da meta e dívida em alta, o ajuste é necessário qualquer que "
                                 "seja o resultado eleitoral."),
    ("Indústria em transição", "A retomada depende dos cortes de juros chegarem ao crédito. CBS e fim do IPI "
                               "exigem preparo das empresas."),
    ("Foco na produtividade", "Com o fim do bônus demográfico, só investimento, inovação e capital humano "
                                  "sustentam o crescimento."),
]


def _titulo(slide, negrito, resto):
    """Título no padrão das edições manuais: início em negrito, resto regular, Montserrat 20 pt azul."""
    p = slide.shapes[0].text_frame.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for txt, b in ((negrito, True), (resto, None)):
        if txt:
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = "Montserrat", Pt(20), b
            r.font.color.rgb = BLUE


def _limpar(slide):
    for sh in list(slide.shapes):
        if sh.name in ("coluna_esq", "coluna_dir"):
            sh._element.getparent().remove(sh._element)


def _caixa(slide, nome, x, y, w, h, partes, alinhamento, ancora=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = nome
    tb.text_frame.vertical_anchor = ancora
    _texto(tb.text_frame, partes, alinhamento)
    return tb


def _seta(slide, x1, x2, y, cor):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    _sem_estilo(ln)
    ln.line.color.rgb = cor
    ln.line.width = Pt(1.5)
    ln.line._get_or_add_ln().append(ln.line._get_or_add_ln().makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def mapa_de_riscos(slide):
    _titulo(slide, "Balanço de riscos", " e o que acompanhar em 2027")
    _limpar(slide)
    cx, pw = 5.0, 2.3  # centro e largura da coluna de temas
    xe, xd = cx - pw / 2, cx + pw / 2
    _caixa(slide, "cab_pior", 0.45, 1.38, xe - 0.55, 0.3, [[("Cenário adverso", 12, True, INK)]], PP_ALIGN.RIGHT)
    _caixa(slide, "cab_tema", xe, 1.38, pw, 0.3, [[("Tema e indicador", 9, False, MUTED)]], PP_ALIGN.CENTER)
    _caixa(slide, "cab_melhor", xd + 0.1, 1.38, 9.55 - xd - 0.1, 0.3, [[("Cenário favorável", 12, True, BLUE)]], PP_ALIGN.LEFT)
    y0, dy, h = 1.82, 0.64, 0.52
    for i, (tema, indicador, pior, melhor) in enumerate(RISCOS):
        y = y0 + i * dy
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(xe), Inches(y), Inches(pw), Inches(h))
        _sem_estilo(pill)
        pill.name = f"tema_{i + 1}"
        pill.adjustments[0] = 0.5
        pill.fill.solid()
        pill.fill.fore_color.rgb = TINTA_AZUL
        pill.line.fill.background()
        pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _texto(pill.text_frame, [[(tema, 11, True, BLUE)], [(indicador, 8, False, MUTED)]])
        pill.text_frame.paragraphs[1].space_before = Pt(1)
        meio = y + h / 2
        _seta(slide, xe - 0.08, xe - 0.38, meio, PISTA)
        _seta(slide, xd + 0.08, xd + 0.38, meio, BLUE)
        _caixa(slide, f"pior_{i + 1}", 0.45, y, xe - 0.5 - 0.45, h, [[(pior, 9.5, False, INK)]], PP_ALIGN.RIGHT)
        _caixa(slide, f"melhor_{i + 1}", xd + 0.5, y, 9.55 - xd - 0.5, h, [[(melhor, 9.5, False, INK)]], PP_ALIGN.LEFT)


def mensagens_finais(slide):
    _limpar(slide)
    yr = 2.4
    for largura, cor, tracejado in ((22, PISTA, False), (1.5, BRANCO, True)):
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(yr), Inches(10), Inches(yr))
        _sem_estilo(ln)
        ln.name = "reta_de_chegada"
        ln.line.color.rgb = cor
        ln.line.width = Pt(largura)
        if tracejado:
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    larg, x0 = 2.3, 0.3
    passo = (9.7 - x0 - larg) / (len(MENSAGENS) - 1)
    for i, (titulo, texto) in enumerate(MENSAGENS):
        x = x0 + i * passo
        cx, d = x + larg / 2, 0.56
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(yr - d / 2), Inches(d), Inches(d))
        _sem_estilo(c)
        c.name = f"mensagem_marco_{i + 1}"
        c.fill.solid()
        c.fill.fore_color.rgb = BLUE
        c.line.color.rgb = BRANCO
        c.line.width = Pt(2.5)
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _texto(c.text_frame, [[(str(i + 1), 15, True, BRANCO)]])
        _caixa(slide, f"mensagem_{i + 1}", x, yr + 0.5, larg, 2.1,
               [[(titulo, 13, True, BLUE)], [(texto, 11, False, INK)]], PP_ALIGN.CENTER, MSO_ANCHOR.TOP)


def _remover_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in lst:
        if prs.slides.get(int(sid.get("id"))) is slide:
            prs.part.drop_rel(sid.get(qn("r:id")))
            lst.remove(sid)
            return


def aplicar(prs):
    por_titulo = {s.shapes[0].text_frame.text: s for s in prs.slides if s.shapes and s.shapes[0].has_text_frame}
    acompanhar = por_titulo["O que acompanhar em 2027"]
    riscos = por_titulo["Balanço de riscos para 2027"]
    finais = por_titulo["Mensagens finais"]
    mapa_de_riscos(riscos)
    mensagens_finais(finais)
    _remover_slide(prs, acompanhar)


def main(deck=DECK):
    bkp = deck.parent / "backup" / f"{deck.stem} (antes da agenda {datetime.now():%d-%m %Hh%M}).pptx"
    bkp.parent.mkdir(exist_ok=True)
    shutil.copy2(deck, bkp)
    prs = Presentation(deck)
    aplicar(prs)
    prs.save(deck)
    print("Backup:", bkp, "\nAtualizado:", deck)


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else DECK))
