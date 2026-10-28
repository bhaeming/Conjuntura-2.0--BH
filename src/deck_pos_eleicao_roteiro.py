"""
Roteiro do deck "Pós-eleição 2027" desenhado como uma estrada: as seis seções são marcos numerados
sobre uma pista sinuosa, com os rótulos alternando acima e abaixo dela.

Aplica no deck já editado à mão (só troca o conteúdo do slide 2, com backup):

    python src/deck_pos_eleicao_roteiro.py
"""
import math
import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs_provisorio" / "Pós-eleição 2027 - Temas econômicos (versão final).pptx"

BLUE, INK, PISTA, BRANCO = RGBColor(0x02, 0x40, 0xCA), RGBColor(0x1E, 0x1E, 0x29), RGBColor(0xB9, 0xBE, 0xC8), RGBColor(0xFF, 0xFF, 0xFF)
ETAPAS = [
    ("Cenário externo", "Crescimento mundial, risco geopolítico, petróleo, câmbio e comércio exterior de SC"),
    ("Atividade, indústria e emprego", "PIB, produção industrial, Santa Catarina e mercado de trabalho"),
    ("Inflação, juros e consumo", "Política monetária, juro real, endividamento e consumo das famílias"),
    ("Contas públicas", "Resultado primário de 2027 e trajetória da dívida"),
    ("Longo prazo", "Demografia, produtividade, investimento, indústria, formação e inovação"),
    ("Agenda 2027", "O que acompanhar, balanço de riscos e mensagens finais"),
]

# geometria (polegadas): marcos alternam entre a crista (rótulo acima) e o vale (rótulo abaixo)
X0, DX, YC, AMP = 1.70, 1.32, 3.30, 0.42
DIAM, LARG_ROTULO = 0.52, 2.55


def _sem_estilo(sh):
    """Remove o p:style herdado (evita sombra/contorno do tema)."""
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def _y(x):
    """Senoide entre os marcos; depois do último, a pista segue reta até a borda (chegada)."""
    x = min(x, X0 + (len(ETAPAS) - 1) * DX)
    return YC - AMP * math.cos(math.pi * (x - X0) / DX)


def _pista(slide, largura_pt, cor, tracejado=False):
    pts = [(i / 100 * 10.0, _y(i / 100 * 10.0)) for i in range(101)]
    ff = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    ff.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    sh = ff.convert_to_shape()
    _sem_estilo(sh)
    sh.fill.background()
    sh.line.color.rgb = cor
    sh.line.width = Pt(largura_pt)
    if tracejado:
        sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = sh.line._get_or_add_ln()
    ln.set("cap", "flat")
    ln.append(ln.makeelement(qn("a:round"), {}))
    return sh


def _texto(tf, partes, alinhamento=PP_ALIGN.CENTER):
    """partes: lista de parágrafos, cada um [(texto, pt, negrito, cor)]."""
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(partes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhamento
        if i:
            p.space_before = Pt(3)
        for txt, pt, negrito, cor in runs:
            r = p.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = "Montserrat", Pt(pt), negrito
            r.font.color.rgb = cor


def estrada(slide, etapas=ETAPAS):
    """Desenha a estrada do roteiro no slide (abaixo do título)."""
    _pista(slide, 22, PISTA)
    _pista(slide, 1.5, BRANCO, tracejado=True)
    for i, (nome, desc) in enumerate(etapas):
        x = X0 + i * DX
        y = _y(x)
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - DIAM / 2), Inches(y - DIAM / 2), Inches(DIAM), Inches(DIAM))
        _sem_estilo(c)
        c.name = f"marco_{i + 1:02d}"
        c.fill.solid()
        c.fill.fore_color.rgb = BLUE
        c.line.color.rgb = BRANCO
        c.line.width = Pt(2.5)
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _texto(c.text_frame, [[(f"{i + 1:02d}", 13, True, BRANCO)]])

        acima = i % 2 == 0
        h = 1.05
        topo = y - DIAM / 2 - 0.12 - h if acima else y + DIAM / 2 + 0.12
        tb = slide.shapes.add_textbox(Inches(x - LARG_ROTULO / 2), Inches(topo), Inches(LARG_ROTULO), Inches(h))
        tb.name = f"etapa_{i + 1:02d}"
        tb.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM if acima else MSO_ANCHOR.TOP
        _texto(tb.text_frame, [[(nome, 13, True, BLUE)], [(desc, 10, False, INK)]])


def main():
    bkp = DECK.parent / "backup" / f"{DECK.stem} (antes da estrada {datetime.now():%d-%m %Hh%M}).pptx"
    bkp.parent.mkdir(exist_ok=True)
    shutil.copy2(DECK, bkp)
    prs = Presentation(DECK)
    s = prs.slides[1]
    assert s.shapes[0].text_frame.text.startswith("Roteiro"), "slide 2 não é o roteiro"
    for sh in list(s.shapes):
        if sh.name in ("coluna_esq", "coluna_dir") or sh.name.startswith(("marco_", "etapa_")) or sh.shape_type == 5:
            sh._element.getparent().remove(sh._element)
    estrada(s)
    prs.save(DECK)
    print("Backup:", bkp, "\nAtualizado:", DECK)


if __name__ == "__main__":
    sys.exit(main())
